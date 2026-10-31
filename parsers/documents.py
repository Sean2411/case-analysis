"""从 zip 中提取 docx、pdf、xlsx、txt、md 的文本。"""

from __future__ import annotations

import io
import zipfile
from dataclasses import dataclass
from pathlib import Path

from config import (
    MAX_FILE_CHARS,
    MAX_MEMBER_BYTES,
    MAX_PDF_PAGES,
    MAX_PROMPT_CHARS,
    MAX_TOTAL_UNCOMPRESSED,
    MAX_XLSX_ROWS,
    MAX_XLSX_SHEETS,
    MAX_ZIP_BYTES,
    MAX_ZIP_FILES,
)
from errors import ParseError

SUPPORTED_SUFFIXES = {".docx", ".pdf", ".xlsx", ".txt", ".md"}


@dataclass
class Document:
    """一个已尝试解析的材料文件。error 非空表示该文件失败，不中断其他文件。"""

    name: str
    suffix: str
    text: str
    char_count: int
    truncated: bool
    error: str | None = None

    @property
    def ok(self) -> bool:
        return bool(self.text.strip())


@dataclass
class ParseOutcome:
    documents: list[Document]
    warnings: list[str]


def decode_text(data: bytes) -> str:
    """依次尝试常见编码。案件材料经常是 GBK。"""

    for encoding in ("utf-8-sig", "utf-8", "gbk"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise ValueError("无法按 UTF-8 或 GBK 解码文本")


def _decode_zip_filename(name: str, flag_bits: int) -> str:
    """还原中文文件名。

    压缩包若没有设置 UTF-8 标志，Python 会把原始字节按 cp437 解开。
    macOS 和较新工具保存的是 UTF-8，旧版 Windows 压缩包则多为 GBK。
    先严格按 UTF-8 解码；不行再试 GBK、Big5。之前一律按 GBK 解码，
    UTF-8 文件名会变成生僻字或 µë° 这类乱码。
    """

    if flag_bits & 0x800:
        return name.replace("\\", "/")
    try:
        raw = name.encode("cp437")
    except UnicodeEncodeError:
        return name.replace("\\", "/")
    for encoding in ("utf-8", "gbk", "big5", "gb18030"):
        try:
            return raw.decode(encoding).replace("\\", "/")
        except UnicodeDecodeError:
            continue
    return name.replace("\\", "/")


def _safe_member_name(info: zipfile.ZipInfo) -> str | None:
    """拒绝绝对路径和 `..`，避免压缩包把文件写到目录之外。只在内存中读，仍要挡住异常文件名。"""

    name = _decode_zip_filename(info.filename, info.flag_bits)
    if not name or name.endswith("/"):
        return None
    parts = Path(name).parts
    if not parts or any(part in {"..", ""} for part in parts):
        return None
    if name.startswith("/") or Path(name).is_absolute():
        return None
    return name


def _is_noise(name: str) -> bool:
    parts = Path(name).parts
    if any(part == "__MACOSX" or part.startswith(".") for part in parts):
        return True
    return Path(name).name.startswith(".")


def extract_text(suffix: str, data: bytes) -> tuple[str, bool]:
    """返回 (文本, 是否因页数或行数上限被截断)。字符截断由调用方统一处理。"""

    suffix = suffix.lower()
    if suffix in {".txt", ".md"}:
        return decode_text(data), False
    if suffix == ".docx":
        return _extract_docx(data), False
    if suffix == ".pdf":
        return _extract_pdf(data)
    if suffix == ".xlsx":
        return _extract_xlsx(data)
    raise ValueError(f"不支持的文件类型：{suffix}")


def _extract_docx(data: bytes) -> str:
    from docx import Document as DocxDocument

    document = DocxDocument(io.BytesIO(data))
    lines = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            if any(cells):
                lines.append("\t".join(cells))
    return "\n".join(lines)


def _extract_pdf(data: bytes) -> tuple[str, bool]:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    if reader.is_encrypted:
        try:
            reader.decrypt("")
        except Exception as exc:  # noqa: BLE001
            raise ValueError("该 PDF 已加密，无法解析") from exc
        if reader.is_encrypted:
            raise ValueError("该 PDF 已加密，无法解析")

    truncated = False
    pages = list(reader.pages)
    if len(pages) > MAX_PDF_PAGES:
        pages = pages[:MAX_PDF_PAGES]
        truncated = True
    chunks: list[str] = []
    for index, page in enumerate(pages, start=1):
        raw = page.extract_text() or ""
        raw = raw.strip()
        if raw:
            chunks.append(f"--- 第{index}页 ---\n{raw}")
    if not chunks:
        raise ValueError("未能从 PDF 提取到文本，可能是扫描件")
    text = "\n\n".join(chunks)
    if truncated:
        text += f"\n[PDF 页数过多，仅提取前 {MAX_PDF_PAGES} 页]"
    return text, truncated


def _extract_xlsx(data: bytes) -> tuple[str, bool]:
    import openpyxl

    workbook = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    truncated = False
    try:
        sheets = workbook.worksheets[:MAX_XLSX_SHEETS]
        if len(workbook.worksheets) > MAX_XLSX_SHEETS:
            truncated = True
        blocks: list[str] = []
        for sheet in sheets:
            rows: list[str] = []
            for row_index, row in enumerate(sheet.iter_rows(values_only=True)):
                if row_index >= MAX_XLSX_ROWS:
                    truncated = True
                    break
                values = ["" if value is None else str(value).strip() for value in row]
                if any(values):
                    rows.append("\t".join(values))
            if rows:
                blocks.append(f"【工作表】{sheet.title}\n" + "\n".join(rows))
        if truncated:
            blocks.append(f"[表格过大，每个工作表最多读取 {MAX_XLSX_ROWS} 行，最多 {MAX_XLSX_SHEETS} 个工作表]")
        return "\n\n".join(blocks), truncated
    finally:
        workbook.close()


def _clip_file_text(text: str, limit: int) -> tuple[str, bool]:
    if len(text) <= limit:
        return text, False
    return text[:limit] + "\n[该文件内容过长，已截断]", True


def parse_zip_bytes(data: bytes, *, max_file_chars: int = MAX_FILE_CHARS) -> ParseOutcome:
    """解析上传的 zip。压缩包本身损坏时抛出 ParseError；单个文件失败写入 Document.error。"""

    if not data:
        raise ParseError("上传的文件为空。")
    if len(data) > MAX_ZIP_BYTES:
        raise ParseError("压缩包过大，请控制在 30MB 以内。")

    try:
        archive_cm = zipfile.ZipFile(io.BytesIO(data))
    except zipfile.BadZipFile as exc:
        raise ParseError("上传的文件不是有效的 zip 压缩包，或压缩包已损坏。") from exc

    with archive_cm as archive:
        return _read_archive(archive, max_file_chars)


def _read_archive(archive: zipfile.ZipFile, max_file_chars: int) -> ParseOutcome:
    try:
        bad_member = archive.testzip()
    except RuntimeError as exc:
        raise ParseError("压缩包可能已加密或无法读取，请上传未加密的 zip。") from exc
    if bad_member:
        raise ParseError(f"压缩包已损坏，无法完整读取：{bad_member}")

    documents: list[Document] = []
    warnings: list[str] = []
    uncompressed = 0
    considered = 0

    for info in archive.infolist():
        if info.is_dir():
            continue
        name = _safe_member_name(info)
        if name is None:
            warnings.append(f"已跳过不安全的压缩包路径：{info.filename}")
            continue
        if _is_noise(name):
            continue

        suffix = Path(name).suffix.lower()
        if suffix not in SUPPORTED_SUFFIXES:
            warnings.append(f"已跳过不支持的文件：{name}")
            continue

        considered += 1
        if considered > MAX_ZIP_FILES:
            warnings.append(f"文件数量超过 {MAX_ZIP_FILES} 个，其余文件已跳过。")
            break
        if info.file_size > MAX_MEMBER_BYTES:
            documents.append(
                Document(
                    name=name,
                    suffix=suffix,
                    text="",
                    char_count=0,
                    truncated=False,
                    error="单个文件超过 8MB，已跳过",
                )
            )
            continue

        uncompressed += info.file_size
        if uncompressed > MAX_TOTAL_UNCOMPRESSED:
            raise ParseError("压缩包解压后体积过大，请拆分后再上传。")

        try:
            payload = archive.read(info)
            text, extractor_truncated = extract_text(suffix, payload)
            original_chars = len(text)
            text, char_truncated = _clip_file_text(text, max_file_chars)
            documents.append(
                Document(
                    name=name,
                    suffix=suffix,
                    text=text,
                    char_count=original_chars,
                    truncated=extractor_truncated or char_truncated,
                    error=None,
                )
            )
        except Exception as exc:  # noqa: BLE001
            # 单个损坏文件不影响同包中的其他材料。
            documents.append(
                Document(
                    name=name,
                    suffix=suffix,
                    text="",
                    char_count=0,
                    truncated=False,
                    error=f"解析失败：{exc}",
                )
            )

    if not documents:
        warnings.append("压缩包中没有可解析的材料（支持 .docx、.pdf、.xlsx、.txt、.md）。")
    return ParseOutcome(documents=documents, warnings=warnings)


def render_for_prompt(documents: list[Document], max_total: int = MAX_PROMPT_CHARS) -> tuple[str, bool]:
    """把材料渲染成提示词。总长度超限时按文件均分截断，避免只保留排在前面的文件。"""

    blocks: list[tuple[str, str]] = []
    for doc in documents:
        if doc.error and not doc.text:
            body = f"（解析失败：{doc.error}）"
        elif not doc.text.strip():
            body = "（文件没有可用文本）"
        else:
            body = doc.text.strip()
        blocks.append((doc.name, body))

    def pack(items: list[tuple[str, str]]) -> str:
        return "\n\n".join(f"【文件】{name}\n{body}" for name, body in items).strip()

    if not blocks:
        return "", False
    if len(pack(blocks)) <= max_total:
        return pack(blocks), False

    share = max(500, max_total // len(blocks) - 80)
    shrunk: list[tuple[str, str]] = []
    for name, body in blocks:
        if len(body) > share:
            body = body[:share] + "\n[该文件在总长度限制下被再次截断]"
        shrunk.append((name, body))
    text = pack(shrunk)
    if len(text) > max_total:
        text = text[:max_total] + "\n[上下文已截断]"
    return text, True
