"""不访问网络的自测。覆盖解析、空法规库、接口失败提示和三部分标题。"""

from __future__ import annotations

import io
import zipfile
from pathlib import Path

import pytest

from agents.llm import chat_complete, explain_llm_exception, locked_temperature
from agents.pipeline import run_analysis
from config import DISCLAIMER, NO_HIT_NOTE, REQUIRED_HEADINGS
from errors import AnalysisError, LLMCallError, ParseError
from parsers.documents import Document, parse_zip_bytes
from rag.retriever import retrieve_laws
from report import has_required_headings
from sample_case import build_sample_zip

ROOT = Path(__file__).resolve().parents[1]


def _assert_three_headings(markdown: str) -> None:
    positions = [markdown.find(heading) for heading in REQUIRED_HEADINGS]
    assert positions[0] >= 0
    assert positions[0] < positions[1] < positions[2]
    assert has_required_headings(markdown)


def _scripted_complete(messages: list[dict[str, str]], temperature: float, *, summary: str = "ok") -> str:
    del temperature
    system = messages[0]["content"]
    user = messages[1]["content"]
    if "阶段：材料摘要" in system:
        return "### 已确认事实\n- [借款合同.txt] 张三向李四借款人民币100000元，期限至2024年12月31日。"
    if "阶段：资料关系" in system:
        return "UNIQUE_DOC_RELATION\n《借款合同.txt》是核心证据。《转账记录.md》用于支撑出借事实。《催收记录.txt》在时间上晚于合同。"
    if "阶段：人物关系" in system:
        return "UNIQUE_PERSON\n张三是借款人，李四是出借人。材料未载明担保人。"
    if "阶段：法律意见" in system:
        return "UNIQUE_OPINION\n材料已确认双方约定了借款金额和期限。是否应支付逾期利息属于法律判断。"
    if "阶段：汇总" in system:
        if summary == "bad":
            return "模型没有按标题输出。"
        extra = ""
        if "未加载法规库" in user:
            extra += "\n未加载法规库，法律意见仅基于模型已有知识\n"
        return (
            "## 一、资料关系\n\n"
            "《借款合同.txt》为核心证据，《转账记录.md》支撑出借，《催收记录.txt》是之后的催告记录。\n\n"
            "## 二、人物/法律关系\n\n"
            "- 张三：借款人\n- 李四：出借人\n- 关系：借款合同项下的债权债务\n\n"
            "## 三、法律意见\n\n"
            "材料已确认的事实：双方在合同中约定借款100000元。\n"
            "法律判断：逾期后果需结合检索到的规则理解，不能反过来改写合同记载。\n"
            f"{extra}"
        )
    raise AssertionError(f"未知阶段：{system[:60]}")


def _minimal_pdf(text: str) -> bytes:
    safe = text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
    stream = f"BT /F1 12 Tf 50 750 Td ({safe}) Tj ET\n".encode("ascii")
    objects = [
        b"1 0 obj\n<< /Type /Catalog /Pages 2 0 R >>\nendobj\n",
        b"2 0 obj\n<< /Type /Pages /Count 1 /Kids [3 0 R] >>\nendobj\n",
        b"3 0 obj\n<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
        b"/Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >>\nendobj\n",
        b"4 0 obj\n<< /Length " + str(len(stream)).encode("ascii") + b" >>\nstream\n" + stream + b"endstream\nendobj\n",
        b"5 0 obj\n<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>\nendobj\n",
    ]
    header = b"%PDF-1.4\n"
    body = b""
    offsets: list[int] = []
    cursor = len(header)
    for obj in objects:
        offsets.append(cursor)
        body += obj
        cursor += len(obj)
    xref = b"xref\n0 6\n0000000000 65535 f \n"
    for offset in offsets:
        xref += f"{offset:010d} 00000 n \n".encode("ascii")
    trailer = f"trailer\n<< /Size 6 /Root 1 0 R >>\nstartxref\n{cursor}\n%%EOF\n".encode("ascii")
    return header + body + xref + trailer


def _build_mixed_zip() -> bytes:
    from docx import Document as DocxDocument
    import openpyxl

    docx_buffer = io.BytesIO()
    document = DocxDocument()
    document.add_paragraph("张三与李四于2024年1月1日签订借款合同，金额100000元。")
    document.save(docx_buffer)

    xlsx_buffer = io.BytesIO()
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "转账"
    sheet.append(["日期", "付款人", "收款人", "金额"])
    sheet.append(["2024-01-02", "李四", "张三", "100000"])
    workbook.save(xlsx_buffer)

    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as archive:
        _zip_write(archive, "合同.docx", docx_buffer.getvalue())
        _zip_write(archive, "statement.pdf", _minimal_pdf("Zhang borrowed 100000 from Li."))
        _zip_write(archive, "流水.xlsx", xlsx_buffer.getvalue())
        _zip_write(archive, "说明.txt", "补充说明：合同原件由李四持有。".encode("utf-8"))
        _zip_write(archive, "备忘.md", "备忘：尚未看到还款凭证。".encode("utf-8"))
        _zip_write(archive, "照片.jpg", b"not-a-real-image")
        _zip_write(archive, "损坏.docx", b"this is not a docx")
    return buffer.getvalue()


def _zip_write(archive: zipfile.ZipFile, name: str, payload: bytes) -> None:
    info = zipfile.ZipInfo(filename=name)
    info.flag_bits |= 0x800
    archive.writestr(info, payload)


def test_parse_supported_formats() -> None:
    outcome = parse_zip_bytes(_build_mixed_zip())
    by_name = {doc.name: doc for doc in outcome.documents}
    assert "张三" in by_name["合同.docx"].text
    assert "100000" in by_name["合同.docx"].text
    assert "Zhang" in by_name["statement.pdf"].text
    assert "100000" in by_name["statement.pdf"].text
    assert "李四" in by_name["流水.xlsx"].text
    assert "100000" in by_name["流水.xlsx"].text
    assert "原件" in by_name["说明.txt"].text
    assert "还款凭证" in by_name["备忘.md"].text
    assert by_name["损坏.docx"].error
    assert by_name["损坏.docx"].text == ""
    assert any("照片.jpg" in warning for warning in outcome.warnings)
    assert all(doc.ok or doc.error for doc in outcome.documents)


def test_parse_sample_zip_and_gbk_truncation() -> None:
    outcome = parse_zip_bytes(build_sample_zip())
    joined = "\n".join(doc.text for doc in outcome.documents)
    assert "张三" in joined and "100000" in joined and "李四" in joined
    assert len(outcome.documents) == 3
    assert all(doc.ok for doc in outcome.documents)

    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as archive:
        info = zipfile.ZipInfo("gbk.txt")
        info.flag_bits |= 0x800
        archive.writestr(info, "张三借款".encode("gbk"))
        archive.writestr("long.txt", "甲" * 120)
        archive.writestr("../secret.txt", "不应被读取")
    parsed = parse_zip_bytes(buffer.getvalue(), max_file_chars=30)
    gbk_doc = next(doc for doc in parsed.documents if doc.name == "gbk.txt")
    long_doc = next(doc for doc in parsed.documents if doc.name == "long.txt")
    assert "张三借款" in gbk_doc.text
    assert long_doc.truncated
    assert long_doc.char_count == 120
    assert "已截断" in long_doc.text
    assert all(".." not in doc.name for doc in parsed.documents)
    assert any("不安全" in warning for warning in parsed.warnings)


def _zip_with_raw_filenames(entries: list[tuple[bytes, bytes]]) -> bytes:
    """生成不带 UTF-8 标志的 zip，文件名按调用方给的原始字节存放。"""

    import struct
    import zlib

    locals_: list[bytes] = []
    centrals: list[bytes] = []
    offset = 0
    for name, content in entries:
        crc = zlib.crc32(content) & 0xFFFFFFFF
        local = struct.pack(
            "<4sHHHHHIIIHH",
            b"PK\x03\x04",
            20,
            0,
            0,
            0,
            0,
            crc,
            len(content),
            len(content),
            len(name),
            0,
        )
        locals_.append(local + name + content)
        central = struct.pack(
            "<4sHHHHHHIIIHHHHHII",
            b"PK\x01\x02",
            20,
            20,
            0,
            0,
            0,
            0,
            crc,
            len(content),
            len(content),
            len(name),
            0,
            0,
            0,
            0,
            0,
            offset,
        )
        centrals.append(central + name)
        offset += len(locals_[-1])
    local_blob = b"".join(locals_)
    central_blob = b"".join(centrals)
    end = struct.pack(
        "<4sHHHHIIH",
        b"PK\x05\x06",
        0,
        0,
        len(entries),
        len(entries),
        len(central_blob),
        len(local_blob),
        0,
    )
    return local_blob + central_blob + end


def test_zip_filenames_prefer_utf8_then_gbk() -> None:
    utf8_name = "workspace/40_关于2024年借款合同.txt"
    gbk_name = "workspace/01_张三起诉状.txt"
    payload = _zip_with_raw_filenames(
        [
            (utf8_name.encode("utf-8"), "借款合同正文".encode("utf-8")),
            (gbk_name.encode("gbk"), "起诉状正文".encode("utf-8")),
        ]
    )
    outcome = parse_zip_bytes(payload)
    names = {doc.name for doc in outcome.documents}
    assert utf8_name in names
    assert gbk_name in names
    by_name = {doc.name: doc for doc in outcome.documents}
    assert "借款合同正文" in by_name[utf8_name].text
    assert "起诉状正文" in by_name[gbk_name].text
    assert all("µ" not in doc.name and "Φ" not in doc.name for doc in outcome.documents)


def test_bad_zip_has_friendly_error() -> None:
    with pytest.raises(ParseError, match="zip"):
        parse_zip_bytes(b"this is not a zip")
    with pytest.raises(ParseError, match="为空"):
        parse_zip_bytes(b"")

    empty = io.BytesIO()
    with zipfile.ZipFile(empty, "w"):
        pass
    outcome = parse_zip_bytes(empty.getvalue())
    assert outcome.documents == []
    assert any("没有可解析" in warning for warning in outcome.warnings)


def test_empty_law_library_degrades(tmp_path: Path) -> None:
    missing = retrieve_laws(tmp_path / "missing", "借款逾期")
    assert missing.loaded is False
    assert missing.hits == []

    empty_dir = tmp_path / "laws"
    empty_dir.mkdir()
    empty = retrieve_laws(empty_dir, "借款逾期")
    assert empty.loaded is False

    outcome = parse_zip_bytes(build_sample_zip())
    result = run_analysis(
        outcome.documents,
        complete=_scripted_complete,
        laws_dir=empty_dir,
        temperature=0.2,
    )
    _assert_three_headings(result.markdown)
    assert result.laws_loaded is False
    assert DISCLAIMER in result.markdown
    assert DISCLAIMER in result.warnings
    assert result.markdown.startswith("## 一、资料关系")


def test_law_retrieval_injects_source(tmp_path: Path) -> None:
    laws = tmp_path / "laws"
    laws.mkdir()
    (laws / "loan.md").write_text(
        "第六百七十六条 借款人未按照约定的期限返还借款的，应当按照约定或者国家有关规定支付逾期利息。\n",
        encoding="utf-8",
    )
    (laws / "unrelated.md").write_text("船舶优先权适用于海难救助报酬。\n", encoding="utf-8")

    found = retrieve_laws(laws, "借款人逾期未返还借款")
    assert found.loaded is True
    assert found.hits
    assert found.hits[0].source == "loan.md"
    assert "逾期利息" in found.hits[0].text

    seen: list[str] = []

    def complete(messages: list[dict[str, str]], temperature: float) -> str:
        if "阶段：法律意见" in messages[0]["content"] or "阶段：汇总" in messages[0]["content"]:
            seen.append(messages[1]["content"])
        return _scripted_complete(messages, temperature)

    outcome = parse_zip_bytes(build_sample_zip())
    result = run_analysis(outcome.documents, complete=complete, laws_dir=laws, temperature=0.2)
    _assert_three_headings(result.markdown)
    assert any("loan.md" in text and "逾期利息" in text for text in seen)
    assert "来源：`loan.md`" in result.markdown
    assert DISCLAIMER not in result.markdown


def test_loaded_library_without_hits_is_not_called_empty(tmp_path: Path) -> None:
    laws = tmp_path / "laws"
    laws.mkdir()
    (laws / "sea.md").write_text("船舶优先权适用于海难救助报酬。\n", encoding="utf-8")
    outcome = parse_zip_bytes(build_sample_zip())
    result = run_analysis(
        outcome.documents,
        complete=_scripted_complete,
        laws_dir=laws,
        temperature=0.2,
    )
    assert result.laws_loaded is True
    assert result.law_hits == []
    assert NO_HIT_NOTE in result.markdown
    assert DISCLAIMER not in result.markdown


def test_output_has_three_headings_on_sample_zip(tmp_path: Path) -> None:
    outcome = parse_zip_bytes(build_sample_zip())
    logs: list[str] = []
    result = run_analysis(
        outcome.documents,
        complete=_scripted_complete,
        laws_dir=tmp_path,
        temperature=0.2,
        progress=logs.append,
    )
    _assert_three_headings(result.markdown)
    assert "资料关系" in result.markdown
    assert "人物/法律关系" in result.markdown
    assert "法律意见" in result.markdown
    assert len(logs) == 6


def test_malformed_summary_still_has_headings(tmp_path: Path) -> None:
    outcome = parse_zip_bytes(build_sample_zip())

    def complete(messages: list[dict[str, str]], temperature: float) -> str:
        return _scripted_complete(messages, temperature, summary="bad")

    result = run_analysis(outcome.documents, complete=complete, laws_dir=tmp_path, temperature=0.2)
    _assert_three_headings(result.markdown)
    assert "UNIQUE_DOC_RELATION" in result.markdown
    assert "UNIQUE_PERSON" in result.markdown
    assert "UNIQUE_OPINION" in result.markdown
    assert DISCLAIMER in result.markdown


def test_kimi_temperature_is_sent_as_one() -> None:
    assert locked_temperature("https://api.moonshot.ai/v1", "kimi-k3.1-pro") == 1.0
    assert locked_temperature("https://api.deepseek.com", "deepseek-chat") is None

    seen: list[float] = []

    class _Message:
        def __init__(self, content: str) -> None:
            self.content = content

    class _Choice:
        def __init__(self, content: str) -> None:
            self.message = _Message(content)

    class _Response:
        def __init__(self, content: str) -> None:
            self.choices = [_Choice(content)]

    class FakeCompletions:
        def create(self, **kwargs: object) -> _Response:
            temperature = float(kwargs["temperature"])  # type: ignore[arg-type]
            seen.append(temperature)
            if temperature != 1:
                raise RuntimeError("Error code: 400 - invalid temperature: only 1 is allowed for this model")
            return _Response("可以分析")

    class FakeChat:
        completions = FakeCompletions()

    class FakeClient:
        chat = FakeChat()

    text = chat_complete(
        [{"role": "user", "content": "hello"}],
        base_url="https://api.moonshot.ai/v1",
        api_key="sk-test",
        model="kimi-k3.1-pro",
        temperature=0.2,
        client=FakeClient(),
    )
    assert text == "可以分析"
    assert seen == [1.0]

    retry = chat_complete(
        [{"role": "user", "content": "hello"}],
        base_url="https://api.example.com",
        api_key="sk-test",
        model="some-model",
        temperature=0.2,
        client=FakeClient(),
    )
    assert retry == "可以分析"
    assert seen[-2:] == [0.2, 1.0]


def test_api_failure_message() -> None:
    class FakeCompletions:
        def create(self, **kwargs: object) -> object:
            del kwargs
            raise RuntimeError("Error code: 401 invalid api key")

    class FakeChat:
        completions = FakeCompletions()

    class FakeClient:
        chat = FakeChat()

    with pytest.raises(LLMCallError, match="API Key 无效"):
        chat_complete(
            [{"role": "user", "content": "hello"}],
            base_url="https://api.example.com",
            api_key="sk-test",
            model="demo-model",
            temperature=0.2,
            client=FakeClient(),
        )

    with pytest.raises(LLMCallError, match="未配置 API Key"):
        chat_complete(
            [{"role": "user", "content": "hello"}],
            base_url="https://api.example.com",
            api_key="  ",
            model="demo-model",
            temperature=0.2,
        )

    assert "请求超时" in explain_llm_exception(TimeoutError("timed out"))
    assert "无法连接到接口" in explain_llm_exception(ConnectionError("connection refused"))
    assert "请求过于频繁" in explain_llm_exception(RuntimeError("rate limit 429"))

    def explode(messages: list[dict[str, str]], temperature: float) -> str:
        del messages, temperature
        raise LLMCallError("模型调用失败：无法连接到接口，请检查 Base URL 与网络。")

    docs = [
        Document(name="借款合同.txt", suffix=".txt", text="张三向李四借款。", char_count=8, truncated=False)
    ]
    with pytest.raises(LLMCallError, match="模型调用失败") as caught:
        run_analysis(docs, complete=explode, laws_dir=ROOT / "missing-laws", temperature=0.2)
    assert "Base URL" in str(caught.value)


def test_refuse_when_no_text() -> None:
    docs = [
        Document(
            name="扫描件.pdf",
            suffix=".pdf",
            text="",
            char_count=0,
            truncated=False,
            error="解析失败：未能从 PDF 提取到文本，可能是扫描件",
        )
    ]

    def unused(messages: list[dict[str, str]], temperature: float) -> str:
        del messages, temperature
        raise AssertionError("没有文本时不应调用模型")

    with pytest.raises(AnalysisError, match="没有成功解析"):
        run_analysis(docs, complete=unused, laws_dir=ROOT, temperature=0.2)


def test_vector_hook_and_keyword_fallback(tmp_path: Path) -> None:
    laws = tmp_path / "laws"
    laws.mkdir()
    (laws / "loan.md").write_text("借款人应当按照约定返还借款并支付利息。\n", encoding="utf-8")
    (laws / "sea.md").write_text("船舶优先权适用于海难救助。\n", encoding="utf-8")

    def embed(text: str) -> list[float]:
        return [1.0 if "借款" in text else 0.0, 1.0 if "船舶" in text else 0.0]

    vector_hits = retrieve_laws(laws, "借款逾期", embed_fn=embed)
    assert vector_hits.hits
    assert vector_hits.hits[0].source == "loan.md"

    def broken(text: str) -> list[float]:
        del text
        raise RuntimeError("embedding down")

    fallback = retrieve_laws(laws, "借款利息", embed_fn=broken)
    assert fallback.hits
    assert fallback.hits[0].source == "loan.md"
    assert any("向量检索失败" in warning for warning in fallback.warnings)


def test_shipped_demo_laws_are_searchable() -> None:
    result = retrieve_laws(ROOT / "data" / "laws", "自然人借款逾期未还，是否支付逾期利息")
    assert result.loaded is True
    assert result.file_count >= 2
    assert any("借款合同" in hit.source and "逾期利息" in hit.text for hit in result.hits)


def test_app_module_imports() -> None:
    import app

    assert callable(app.main)
